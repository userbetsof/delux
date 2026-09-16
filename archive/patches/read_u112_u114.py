import openpyxl

wb = openpyxl.load_workbook("DeluX_IncoProrok_mvp.xlsm", read_only=True, data_only=False)
ws = wb["Расчеты"]
print("Formula in U112:", ws["U112"].value)
print("Formula in U114:", ws["U114"].value)
wb.close()
