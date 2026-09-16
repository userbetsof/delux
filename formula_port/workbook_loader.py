"""Загрузка формул из кэша (или .xlsm при сборке кэша). Runtime не требует Excel."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Dict, Iterable, List

ROOT = Path(__file__).resolve().parent.parent
CACHE_PATH = ROOT / "data" / "formulas_cache.json"
CONST_CACHE_PATH = ROOT / "data" / "constants_cache.json"


def has_formula_cache() -> bool:
    return CACHE_PATH.is_file()


def _constants_from_lookups(wanted: set) -> Dict[str, Dict[str, object]]:
    """Неформульные константы из data/lookups.db — без .xlsm."""
    from formula_port.static_cache import StaticCache

    cache = StaticCache()
    if not cache.exists():
        return {}
    by_sheet: Dict[str, Dict[str, object]] = {}
    for key, val in cache.load_all().items():
        if "!" not in key:
            continue
        sheet, addr = key.split("!", 1)
        if wanted and sheet not in wanted:
            continue
        by_sheet.setdefault(sheet, {})[addr] = val
    return by_sheet


def load_sheet_formulas_from_workbook(
    workbook_path: str, wanted: set
) -> Dict[str, Dict[str, str]]:
    import openpyxl

    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=False)
    out: Dict[str, Dict[str, str]] = {}
    for name in wb.sheetnames:
        if name not in wanted:
            continue
        ws = wb[name]
        formulas = {}
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if val is not None and str(val).startswith("="):
                    formulas[cell.coordinate] = str(val)
        out[name] = formulas
    wb.close()
    return out


def load_sheet_formulas(workbook_path: str, sheet_names: Iterable[str]) -> Dict[str, Dict[str, str]]:
    wanted = set(sheet_names)
    if CACHE_PATH.exists():
        with open(CACHE_PATH, encoding="utf-8") as f:
            cached = json.load(f)
        sheets = cached.get("sheets", {})
        if sheets:
            return {k: v for k, v in sheets.items() if (not wanted or k in wanted)}
    if workbook_path and Path(workbook_path).is_file():
        return load_sheet_formulas_from_workbook(workbook_path, wanted)
    raise FileNotFoundError(
        f"Нет кэша формул ({CACHE_PATH}) и нет .xlsm. "
        "Положите data/formulas_cache.json или укажите CALC_WORKBOOK."
    )


def _load_constants_from_workbook(
    workbook_path: str, wanted: set
) -> Dict[str, Dict[str, object]]:
    import openpyxl

    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=False)
    out: Dict[str, Dict[str, object]] = {}
    for name in wb.sheetnames:
        if name not in wanted:
            continue
        ws = wb[name]
        values = {}
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if val is None:
                    continue
                if isinstance(val, str) and val.startswith("="):
                    continue
                values[cell.coordinate] = val
        out[name] = values
    wb.close()
    return out


def load_sheet_constants(
    workbook_path: str, sheet_names: Iterable[str]
) -> Dict[str, Dict[str, object]]:
    """Неформульные значения на листах (стартовые константы)."""
    wanted = set(sheet_names)
    if CONST_CACHE_PATH.exists():
        with open(CONST_CACHE_PATH, encoding="utf-8") as f:
            cached = json.load(f)
        sheets = cached.get("sheets", {})
        if sheets:
            return {k: v for k, v in sheets.items() if (not wanted or k in wanted)}
    from_lookups = _constants_from_lookups(wanted)
    if from_lookups:
        return from_lookups
    if workbook_path and Path(workbook_path).is_file():
        data = _load_constants_from_workbook(workbook_path, wanted)
        CONST_CACHE_PATH.parent.mkdir(parents=True, exist_ok=True)
        with open(CONST_CACHE_PATH, "w", encoding="utf-8") as f:
            json.dump({"sheets": data}, f, ensure_ascii=False)
        return data
    return {}


def populate_formula_store(workbook_path: str, sheet_names: List[str]):
    from formula_port.cell_store import CellStore

    store = CellStore()
    formulas = load_sheet_formulas(workbook_path, sheet_names)
    for sheet, mapping in formulas.items():
        store.bulk_set_formulas(sheet, mapping)
    return store
