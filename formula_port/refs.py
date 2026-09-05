"""Парсинг ссылок Excel (ячейки, диапазоны, листы)."""

from __future__ import annotations

import re
from dataclasses import dataclass
from typing import Iterator, List, Optional, Tuple

from openpyxl.utils import column_index_from_string, get_column_letter

CELL_RE = re.compile(
    r"(?:'([^']+)'|([A-Za-z\u0410-\u042f\u0430-\u044f\u0401\u04510-9 _]+))!"
    r"(\$?)([A-Z]{1,3})(\$?)(\d+)"
    r"(?::(\$?)([A-Z]{1,3})(\$?)(\d+))?",
    re.IGNORECASE,
)
BARE_CELL_RE = re.compile(r"^(\$?)([A-Z]{1,3})(\$?)(\d+)$", re.IGNORECASE)
SHEET_LOCAL_RANGE_RE = re.compile(
    r"(\$?)([A-Z]{1,3})(\$?)(\d+):(\$?)([A-Z]{1,3})(\$?)(\d+)",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class CellRef:
    sheet: str
    row: int
    col: int

    def key(self) -> str:
        return f"{self.sheet}!{get_column_letter(self.col)}{self.row}"

    def a1(self) -> str:
        return f"{get_column_letter(self.col)}{self.row}"


def col_row(addr: str) -> Tuple[int, int]:
    addr = addr.replace("$", "").upper()
    m = BARE_CELL_RE.match(addr)
    if not m:
        raise ValueError(f"Bad cell address: {addr}")
    return column_index_from_string(m.group(2)), int(m.group(4))


def iter_range(sheet: str, start_col: int, start_row: int, end_col: int, end_row: int) -> Iterator[CellRef]:
    for r in range(start_row, end_row + 1):
        for c in range(start_col, end_col + 1):
            yield CellRef(sheet, r, c)


def expand_ref(sheet: str, start_a1: str, end_a1: Optional[str] = None) -> List[CellRef]:
    sc, sr = col_row(start_a1)
    if end_a1:
        ec, er = col_row(end_a1)
    else:
        ec, er = sc, sr
    if sc > ec:
        sc, ec = ec, sc
    if sr > er:
        sr, er = er, sr
    return list(iter_range(sheet, sc, sr, ec, er))


def parse_formula_refs(formula: str, default_sheet: str = "") -> List[CellRef]:
    """Извлекает ссылки на ячейки из формулы Excel."""
    if not formula or not str(formula).startswith("="):
        return []
    text = str(formula)
    refs: List[CellRef] = []
    seen = set()

    covered = []
    for m in CELL_RE.finditer(text):
        covered.append((m.start(), m.end()))
        sheet = m.group(1) or m.group(2) or default_sheet
        c1 = column_index_from_string(m.group(4).upper())
        r1 = int(m.group(6))
        if m.group(7) is not None:
            c2 = column_index_from_string(m.group(8).upper())
            r2 = int(m.group(10))
            for ref in iter_range(sheet, c1, r1, c2, r2):
                if ref.key() not in seen:
                    seen.add(ref.key())
                    refs.append(ref)
        else:
            ref = CellRef(sheet, r1, c1)
            if ref.key() not in seen:
                seen.add(ref.key())
                refs.append(ref)

    if default_sheet:
        def _covered(start: int, end: int) -> bool:
            return any(s <= start and end <= e for s, e in covered)

        for m in SHEET_LOCAL_RANGE_RE.finditer(text):
            if _covered(m.start(), m.end()):
                continue
            c1 = column_index_from_string(m.group(2).upper())
            r1 = int(m.group(4))
            c2 = column_index_from_string(m.group(6).upper())
            r2 = int(m.group(8))
            for ref in iter_range(default_sheet, c1, r1, c2, r2):
                if ref.key() not in seen:
                    seen.add(ref.key())
                    refs.append(ref)

        for m in re.finditer(r"(?<![A-Z0-9_!])(\$?[A-Z]{1,3}\$?\d+)(?![A-Z0-9_])", text, re.I):
            if _covered(m.start(), m.end()):
                continue
            token = m.group(1)
            if ":" in token:
                continue
            try:
                c, r = col_row(token)
            except ValueError:
                continue
            ref = CellRef(default_sheet, r, c)
            if ref.key() not in seen:
                seen.add(ref.key())
                refs.append(ref)

    return refs
