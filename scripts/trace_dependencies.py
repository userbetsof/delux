"""BFS граф зависимостей от выходных ячеек Анализ → data/formula_graph.json."""

from __future__ import annotations

import json
import sys
from collections import deque
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from formula_port.refs import CellRef, parse_formula_refs  # noqa: E402
from settings import DeluxSettings  # noqa: E402
from sheets_config import (  # noqa: E402
    DYNAMIC_SHEETS,
    SHEET_ANALYSIS,
    SHEET_AWAY,
    SHEET_AWAY_AWAY,
    SHEET_AWAY_HOME,
    SHEET_CALC,
    SHEET_FORM,
    SHEET_H2H,
    SHEET_H2H_HOME,
    SHEET_HOME,
    SHEET_HOME_AWAY,
    SHEET_HOME_HOME,
    SHEET_LIST,
    SHEET_ODDS,
    SHEET_START,
    SHEET_TABLE,
    STATIC_SHEETS,
)

OUTPUT_CELLS = [
    (SHEET_ANALYSIS, "S13"),
    (SHEET_ANALYSIS, "O10"),
    (SHEET_ANALYSIS, "X24"),
    (SHEET_ANALYSIS, "E2"),
    (SHEET_ANALYSIS, "K2"),
    (SHEET_ANALYSIS, "E3"),
    (SHEET_ANALYSIS, "K3"),
]
for idx in range(10):
    row = 8 + idx * 2
    for col in ("H", "I", "J", "K", "L"):
        OUTPUT_CELLS.append((SHEET_ANALYSIS, f"{col}{row}"))
    OUTPUT_CELLS.append((SHEET_ANALYSIS, f"J{row + 1}"))
for col_n in list(range(22, 28)) + list(range(35, 41)):
    from openpyxl.utils import get_column_letter

    letter = get_column_letter(col_n)
    OUTPUT_CELLS.append((SHEET_ANALYSIS, f"{letter}10"))
    OUTPUT_CELLS.append((SHEET_ANALYSIS, f"{letter}11"))

INPUT_SHEETS = {
    SHEET_H2H,
    SHEET_H2H_HOME,
    SHEET_HOME,
    SHEET_HOME_HOME,
    SHEET_HOME_AWAY,
    SHEET_AWAY,
    SHEET_AWAY_HOME,
    SHEET_AWAY_AWAY,
    SHEET_TABLE,
    SHEET_FORM,
    SHEET_ODDS,
    SHEET_LIST,
    SHEET_START,
}


def load_workbook_formulas(path: str) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    store: dict = {}
    for name in wb.sheetnames:
        ws = wb[name]
        sheet_formulas = {}
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if val is not None and str(val).startswith("="):
                    sheet_formulas[cell.coordinate] = str(val)
        store[name] = sheet_formulas
    wb.close()
    return store


def classify_sheet(sheet: str) -> str:
    if sheet in INPUT_SHEETS:
        return "input"
    if sheet in DYNAMIC_SHEETS:
        return "dynamic"
    if sheet in STATIC_SHEETS:
        return "static"
    if sheet in {SHEET_HOME, SHEET_AWAY, SHEET_HOME_HOME, SHEET_HOME_AWAY, SHEET_AWAY_HOME, SHEET_AWAY_AWAY}:
        return "venue"
    return "other"


def trace(workbook_path: str) -> dict:
    formulas = load_workbook_formulas(workbook_path)
    nodes: dict = {}
    edges: list = []
    queue: deque = deque()

    for sheet, addr in OUTPUT_CELLS:
        key = f"{sheet}!{addr}"
        nodes[key] = {"sheet": sheet, "cell": addr, "role": "output", "sheet_class": classify_sheet(sheet)}
        queue.append((sheet, addr))

    while queue:
        sheet, addr = queue.popleft()
        key = f"{sheet}!{addr}"
        formula = formulas.get(sheet, {}).get(addr)
        if formula:
            nodes[key]["formula"] = formula
        else:
            nodes[key].setdefault("formula", None)

        if not formula:
            continue

        for ref in parse_formula_refs(formula, default_sheet=sheet):
            ref_key = ref.key()
            if ref_key not in nodes:
                nodes[ref_key] = {
                    "sheet": ref.sheet,
                    "cell": ref.a1(),
                    "sheet_class": classify_sheet(ref.sheet),
                }
                queue.append((ref.sheet, ref.a1()))
            edges.append({"from": ref_key, "to": key})

    by_class: dict = {}
    for node in nodes.values():
        cls = node.get("sheet_class", "other")
        by_class.setdefault(cls, []).append(f"{node['sheet']}!{node['cell']}")

    static_cells = [k for k, v in nodes.items() if v.get("sheet_class") == "static"]
    dynamic_cells = [k for k, v in nodes.items() if v.get("sheet_class") == "dynamic"]
    venue_cells = [k for k, v in nodes.items() if v.get("sheet_class") in ("venue", "input")]

    return {
        "workbook": workbook_path,
        "output_count": len(OUTPUT_CELLS),
        "node_count": len(nodes),
        "edge_count": len(edges),
        "by_class_counts": {k: len(v) for k, v in by_class.items()},
        "nodes": nodes,
        "edges": edges,
        "static_cells": sorted(static_cells),
        "dynamic_cells": sorted(dynamic_cells),
        "venue_cells": sorted(venue_cells),
    }


def main():
    cfg = DeluxSettings.from_env()
    workbook = cfg.resolve_calc_workbook()
    graph = trace(workbook)
    out_dir = ROOT / "data"
    out_dir.mkdir(exist_ok=True)
    out_path = out_dir / "formula_graph.json"
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(graph, f, ensure_ascii=False, indent=2)
    print(f"[+] Graph: {graph['node_count']} nodes, {graph['edge_count']} edges")
    print(f"    classes: {graph['by_class_counts']}")
    print(f"    saved: {out_path}")


if __name__ == "__main__":
    main()
