import openpyxl

wb = openpyxl.load_workbook("DeluX_IncoProrok_mvp.xlsm", read_only=True, data_only=True)
sheet = wb["Список"]
leagues = set()
for row in range(2, sheet.max_row + 1):
    val = sheet.cell(row, 1).value
    if val:
        leagues.add(str(val).strip())

print("Leagues in 'Список' sheet:")
for l in sorted(leagues):
    print(f"  - {l}")
wb.close()
