import openpyxl

wb = openpyxl.load_workbook("DeluX_IncoProrok_mvp.xlsm", read_only=True, data_only=True)
sheet = wb["achive"]
found = False
for row in range(1, sheet.max_row + 1):
    val = sheet.cell(row, 1).value
    if val == "IkY7Oqwp":
        print(f"Match IkY7Oqwp found in achive sheet at row {row}")
        found = True
        break
if not found:
    print("Match IkY7Oqwp not found in achive sheet.")
wb.close()
